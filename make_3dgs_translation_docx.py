# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\Users\Administrator\Documents\my 3dgs\3DGS_initial_paper_Chinese_translation.docx")


def set_run_font(run, size=10.5, bold=False, italic=False, color=None, name="Microsoft YaHei"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size=9.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell(cell, h, bold=True)
        shade_cell(cell, "E8EEF5")
        if widths:
            cell.width = Cm(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell(cells[i], value)
            if widths:
                cells[i].width = Cm(widths[i])
    doc.add_paragraph("")
    return table


def add_para(doc, text, style=None, size=10.5, bold=False, italic=False, color=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(item)
        set_run_font(run, size=10)


def setup_doc():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    styles = doc.styles
    for name in ["Normal", "Heading 1", "Heading 2", "Heading 3", "List Bullet", "List Number"]:
        styles[name].font.name = "Microsoft YaHei"
        styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles["Normal"].font.size = Pt(10.5)
    styles["Heading 1"].font.size = Pt(16)
    styles["Heading 1"].font.color.rgb = RGBColor(46, 116, 181)
    styles["Heading 2"].font.size = Pt(13)
    styles["Heading 2"].font.color.rgb = RGBColor(46, 116, 181)
    styles["Heading 3"].font.size = Pt(12)
    styles["Heading 3"].font.color.rgb = RGBColor(31, 77, 120)
    return doc


def add_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("3D Gaussian Splatting for Real-Time Radiance Field Rendering")
    set_run_font(run, size=18, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("实时辐射场渲染的三维高斯泼溅")
    set_run_font(run, size=16, bold=True, color=(31, 77, 120))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("非官方中文译文整理版")
    set_run_font(run, size=10.5, italic=True, color=(90, 90, 90))

    add_para(
        doc,
        "原文：Bernhard Kerbl, Georgios Kopanas, Thomas Leimkühler, George Drettakis. "
        "ACM Transactions on Graphics, SIGGRAPH 2023。本文档基于用户提供的 2308.04079v1.pdf 制作，"
        "为便于学习，将双栏论文重排为单栏中文，并保留关键英文术语、公式和图表编号。译文非官方版本。",
        size=9.5,
        color=(80, 80, 80),
    )


def main():
    doc = setup_doc()
    add_title(doc)

    doc.add_heading("术语约定", level=1)
    add_table(
        doc,
        ["英文术语", "中文译法", "说明"],
        [
            ("radiance field", "辐射场", "描述空间位置和观察方向对应的颜色/密度等信息。"),
            ("novel view synthesis", "新视角合成", "从未见过的相机位姿生成图像。"),
            ("3D Gaussian", "三维高斯", "带中心、协方差、颜色和不透明度的可优化基元。"),
            ("splatting", "泼溅 / splatting", "把三维基元投影成屏幕上的二维斑点并混合。"),
            ("opacity", "不透明度", "控制 Gaussian 对像素贡献和遮挡的强度。"),
            ("spherical harmonics", "球谐函数 / SH", "用于表示视角相关颜色。"),
            ("rasterizer", "光栅化器", "把 Gaussian 投影、排序并合成为图像的渲染器。"),
        ],
        [3.6, 3.6, 8.2],
    )

    doc.add_heading("摘要", level=1)
    add_para(
        doc,
        "近年来，辐射场方法极大推动了基于多张照片或视频的新视角合成。然而，要获得高视觉质量，"
        "通常仍依赖训练和渲染代价高昂的神经网络；而近期一些更快的方法又不可避免地在速度和质量之间折中。"
        "对于完整的、无界的真实场景，而不是孤立物体，在 1080p 分辨率下实现实时显示，过去方法仍难以同时满足。"
    )
    add_para(
        doc,
        "本文提出三个关键要素，使方法能够在保持有竞争力训练时间的同时达到先进视觉质量，并重要地实现 "
        "1080p 下高质量实时（大于等于 30 FPS）新视角合成。第一，从相机标定过程中产生的稀疏点出发，"
        "用三维高斯表示场景；这种表示保留了连续体积辐射场适合优化的性质，同时避免在空空间中做不必要计算。"
        "第二，对三维高斯的各向异性协方差进行优化，并与自适应密度控制交替进行，从而得到对场景的紧凑而准确的表示。"
        "第三，提出快速、可微、可见性感知的 splatting 渲染器，使训练和实时显示都能高效进行。"
    )
    add_para(
        doc,
        "实验表明，该方法在多个真实数据集上达到或超过当时的先进质量，同时训练时间与最快的先前方法相当，"
        "并显著提高新视角渲染速度。"
    )

    doc.add_heading("1 引言", level=1)
    add_para(
        doc,
        "网格和点是最常见的三维场景表示，因为它们是显式表示，并且非常适合基于 GPU / CUDA 的快速光栅化。"
        "然而，真实照片采集的场景往往具有复杂几何、薄结构、反射、高频纹理和不完整观测，直接恢复干净网格并不容易。"
        "另一方面，NeRF 类方法用连续辐射场表达场景，能够产生高质量图像，但其训练和渲染通常需要大量网络查询或沿光线采样，"
        "难以在完整场景和高分辨率下实时运行。"
    )
    add_para(
        doc,
        "本文的核心思想是：采用一种显式但连续的场景表示，即一组三维高斯基元。每个高斯具有位置、形状、透明度和视角相关颜色。"
        "渲染时，不沿每条光线密集采样，而是把这些三维高斯投影到屏幕上，形成二维椭圆 splat，再按深度进行 alpha 合成。"
        "因此，该方法将传统显式图形学表示的高效性，与辐射场方法的可微优化和照片级合成能力结合起来。"
    )
    add_para(
        doc,
        "图 1 展示了本文方法与 Instant-NGP、Plenoxels、Mip-NeRF360 等方法的比较。作者强调：在与最快方法相近的训练时间内，"
        "本文方法能达到类似质量；若训练更久，则能达到当时先进质量，并保持实时渲染速度。"
    )
    add_para(
        doc,
        "本文的贡献可以概括为三点："
    )
    add_bullets(
        doc,
        [
            "提出以三维高斯为基元的显式辐射场表示，适合真实、完整、无界场景的新视角合成。",
            "提出适用于三维高斯的优化与自适应密度控制策略，包括增密、分裂、复制和剪枝。",
            "提出快速、可微、可见性感知的 splatting 渲染器，使 1080p 实时渲染成为可能。",
        ],
    )

    doc.add_heading("2 相关工作", level=1)
    add_para(
        doc,
        "相关工作主要包括三个方向：新视角合成和辐射场、基于点的渲染与重建、以及可微渲染。"
        "NeRF 及其后续工作通过神经网络表示连续辐射场，显著提高了新视角合成质量。"
        "Mip-NeRF360 进一步处理无界场景和抗锯齿问题，但训练和渲染代价仍然很高。"
        "Instant-NGP 通过多分辨率哈希网格加速训练和查询，Plenoxels 则使用显式体素网格来避免大网络，但这些方法仍存在速度、内存或质量上的权衡。"
    )
    add_para(
        doc,
        "点渲染和 splatting 有较长历史。点基元适合快速渲染，但普通点缺少连续面积和各向异性形状，容易产生孔洞或不稳定伪影。"
        "早期的椭圆加权平均 splatting、surfel 渲染等工作已经使用屏幕空间椭圆足迹来改善点渲染质量。"
        "本文继承这种思想，但把可优化的三维高斯作为场景表示，并使其参与可微训练。"
    )
    add_para(
        doc,
        "可微渲染方法试图让图像误差反向传播到三维场景参数。许多方法依赖网格、体素或神经隐式场。"
        "本文的不同之处在于：它设计了一个面向三维高斯的可微光栅化器，并把高斯的形状、透明度、颜色和密度控制都纳入训练。"
    )

    doc.add_heading("3 方法概览", level=1)
    add_para(
        doc,
        "方法输入是一组经过 SfM / COLMAP 处理的照片，包括相机内外参和稀疏点云。"
        "算法首先把稀疏点云中的每个点初始化为一个三维高斯。随后通过多视角图像监督，优化每个高斯的位置、协方差、颜色和不透明度。"
        "训练过程中还会周期性执行自适应密度控制：在误差敏感区域增加高斯，在无用区域删除高斯。"
    )
    add_para(
        doc,
        "一次训练迭代可以理解为：随机选一张训练图像及其相机；用当前三维高斯集合渲染该视角；"
        "把渲染图与真实图比较得到损失；通过可微渲染器反向传播梯度；更新高斯参数；再根据梯度和尺度决定是否增密或剪枝。"
    )

    doc.add_heading("4 可微三维高斯泼溅", level=1)
    doc.add_heading("4.1 三维高斯表示", level=2)
    add_para(
        doc,
        "每个高斯由中心位置 μ 和协方差矩阵 Σ 定义，其空间影响可写为："
    )
    add_para(doc, "G(x) = exp(-1/2 (x - μ)^T Σ^{-1} (x - μ))", italic=True, color=(31, 77, 120))
    add_para(
        doc,
        "直接优化任意协方差矩阵可能导致矩阵不合法。为保证协方差正定或半正定，论文把协方差分解为旋转和缩放："
    )
    add_para(doc, "Σ = R S S^T R^T", italic=True, color=(31, 77, 120))
    add_para(
        doc,
        "其中 S 是对角缩放矩阵，控制三个主轴方向的尺度；R 是旋转矩阵，控制椭球朝向。"
        "直观地说，先把单位球沿三个方向缩放成椭球，再把椭球旋转到场景中的合适方向。"
        "在实现中，旋转通常用四元数存储并归一化，尺度使用指数激活保证为正。"
    )

    doc.add_heading("4.2 投影到图像平面", level=2)
    add_para(
        doc,
        "渲染时，需要把三维高斯投影成屏幕空间的二维高斯。高斯中心按照普通相机投影变成二维位置。"
        "对于协方差，由于透视投影是非线性的，论文在高斯中心附近使用一阶泰勒近似，即用投影函数的雅可比矩阵 J 近似局部变换。"
        "二维协方差写为："
    )
    add_para(doc, "Σ' = J W Σ W^T J^T", italic=True, color=(31, 77, 120))
    add_para(
        doc,
        "其中 W 是世界坐标到相机坐标的变换，J 是投影在当前中心附近的雅可比矩阵。"
        "这个公式本质上是协方差在线性变换下的传播规律：Σ_y = A Σ_x A^T。"
        "因此，一个三维椭球从某个相机看过去，就变成屏幕上的一个二维椭圆 splat。"
    )

    doc.add_heading("4.3 颜色和透明度合成", level=2)
    add_para(
        doc,
        "每个高斯还带有不透明度 α 和颜色 c。颜色不是简单固定 RGB，而是用球谐函数表示，以便描述一定程度的视角相关外观。"
        "对于一个像素 p，高斯的有效 alpha 等于高斯自身 opacity 与该像素处二维高斯权重的乘积。"
    )
    add_para(
        doc,
        "当多个高斯覆盖同一个像素时，需要按深度从近到远进行 alpha 合成。像素颜色为："
    )
    add_para(doc, "C = Σ_i c_i α_i Π_{j<i} (1 - α_j)", italic=True, color=(31, 77, 120))
    add_para(
        doc,
        "这里 Π_{j<i}(1-α_j) 表示前面更近的高斯遮挡后，后方第 i 个高斯仍然可见的比例。"
        "这与体渲染中的透射率概念一致，但计算方式变成了屏幕空间 splatting，而不是沿每条光线大量采样。"
    )

    doc.add_heading("5 带自适应密度控制的优化", level=1)
    add_para(
        doc,
        "仅从 COLMAP 稀疏点云初始化的高斯远远不够覆盖完整场景。论文因此在训练中交替进行连续参数优化和密度控制。"
        "连续优化负责移动高斯、改变形状、调整颜色和透明度；密度控制负责在需要更多表达能力的地方增加高斯，在无用区域删除高斯。"
    )
    doc.add_heading("5.1 损失函数", level=2)
    add_para(
        doc,
        "渲染图像与真实训练图像之间的主要损失由 L1 和 D-SSIM 组成："
    )
    add_para(doc, "L = (1 - λ) L1 + λ L_D-SSIM", italic=True, color=(31, 77, 120))
    add_para(
        doc,
        "L1 负责像素颜色接近，D-SSIM 负责局部结构和视觉相似性。论文中常用 λ = 0.2。"
        "损失在二维图像上计算，但通过可微渲染器反传到三维高斯的位置、尺度、旋转、颜色和不透明度。"
    )
    doc.add_heading("5.2 增密：clone 与 split", level=2)
    add_para(
        doc,
        "论文使用视空间位置梯度判断哪些高斯所在区域表达能力不足。如果一个高斯的屏幕位置梯度较大，说明它附近的图像误差对其位置很敏感，"
        "通常意味着边缘、纹理、遮挡交界或细节没有被充分表示。"
    )
    add_para(
        doc,
        "若该高斯尺度较小且梯度大，则执行 clone：复制一个相似高斯，用更多局部基元表达细节。"
        "若该高斯尺度较大且梯度大，则执行 split：在原高斯分布附近采样多个新高斯，并把尺度缩小，同时删除原来的大高斯。"
        "这样大而粗糙的高斯会被拆成更精细的局部表示。"
    )
    doc.add_heading("5.3 剪枝与 opacity reset", level=2)
    add_para(
        doc,
        "为了防止高斯数量无限增长，训练中会删除 opacity 很低、屏幕空间半径过大或世界空间尺度异常大的高斯。"
        "opacity 很低说明该高斯几乎没有贡献；半径或尺度过大则可能造成大面积模糊或错误遮挡。"
    )
    add_para(
        doc,
        "此外，训练中会周期性重置 opacity，使部分高斯重新变得透明一些。"
        "这样可以避免错误高斯过早变得不透明并挡住后方结构，从而让更多高斯获得有效梯度，便于后续纠正和剪枝。"
    )

    doc.add_heading("6 用于高斯的快速可微光栅化器", level=1)
    add_para(
        doc,
        "为了达到实时渲染，论文设计了专门的可微光栅化器。它不是对每个像素遍历所有高斯，而是把屏幕划分成 tile。"
        "每个三维高斯投影后只会覆盖有限的二维区域，渲染器只把它分配给相关 tile。随后在每个 tile 内按深度排序，并对像素执行前向 alpha 合成。"
    )
    add_para(
        doc,
        "这种设计减少了无关计算，并适合 GPU 并行。当前方高斯已经使透射率很低时，后方高斯对像素贡献可以忽略，渲染器还可以提前终止。"
        "因此，相比 NeRF 中沿每条光线采样大量点并查询网络的方式，3DGS 只处理实际存在的显式高斯基元，避免了空空间采样。"
    )
    add_para(
        doc,
        "反向传播方面，光栅化器需要把图像损失对像素颜色的梯度传回高斯参数，包括二维投影位置、协方差、颜色和透明度。"
        "论文附录给出了梯度计算的细节。实际代码中，这部分由自定义 CUDA 扩展实现，并通过 PyTorch autograd 接入训练流程。"
    )

    doc.add_heading("7 实现、结果与评估", level=1)
    add_para(
        doc,
        "实验部分比较了本文方法与 Instant-NGP、Plenoxels、Mip-NeRF360 等代表性方法。评价指标包括 PSNR、SSIM、LPIPS、训练时间和渲染帧率。"
        "PSNR 衡量像素误差，SSIM 衡量结构相似性，LPIPS 更接近感知质量，FPS 衡量实时渲染能力。"
    )
    add_para(
        doc,
        "论文报告显示，在多个真实场景数据集上，3DGS 能以远高于 NeRF 类方法的速度进行渲染，同时达到相当甚至更好的图像质量。"
        "尤其在 1080p 分辨率下，本文方法能够实现交互式或实时显示，这是当时许多高质量辐射场方法难以做到的。"
    )
    add_para(
        doc,
        "作者还进行了消融实验，验证各个组件的重要性：各向异性协方差、自适应密度控制、可见性感知排序、球谐颜色和专门光栅化器都会影响最终质量或速度。"
        "总体结论是，3DGS 的成功不只来自使用高斯这一表示，而是来自表示、优化策略和渲染器三者的共同设计。"
    )
    add_table(
        doc,
        ["评估维度", "含义", "对 3DGS 的意义"],
        [
            ("PSNR", "像素级误差，越高越好。", "说明渲染图与真实图颜色更接近。"),
            ("SSIM", "局部结构相似性，越高越好。", "说明边缘、纹理和结构更自然。"),
            ("LPIPS", "感知距离，越低越好。", "更接近人眼观感。"),
            ("FPS", "每秒帧数，越高越好。", "体现实时新视角合成能力。"),
            ("Training time", "训练耗时。", "体现从数据到可用场景表示的效率。"),
        ],
        [3.0, 5.0, 7.0],
    )

    doc.add_heading("8 讨论与局限", level=1)
    add_para(
        doc,
        "尽管 3DGS 能产生高质量图像，但其几何并不一定等同于真实物理表面。"
        "由于监督主要来自二维图像，只要从训练视角和附近新视角渲染正确，损失就会降低；这并不保证每个高斯都准确落在真实表面上。"
        "因此，提取高质量 mesh、处理薄结构或获得严格几何仍然具有挑战。"
    )
    add_para(
        doc,
        "方法还依赖 COLMAP 等 SfM 系统提供相机位姿和稀疏点云。如果输入照片纹理弱、曝光变化大、视角重叠不足，或者场景中有大量动态物体，"
        "SfM 可能失败，从而影响 3DGS 初始化。"
    )
    add_para(
        doc,
        "初始论文主要针对静态场景。对于动态人物、树叶、水面或光照变化，静态高斯集合难以完整表达。"
        "此外，强反光、透明和折射现象也超出了简单 SH 视角相关颜色的表达能力。"
        "这些问题后来催生了动态 3DGS、稀疏视角 3DGS、压缩 3DGS、语义/可编辑 Gaussian 等大量后续研究。"
    )

    doc.add_heading("附录 A：梯度计算要点", level=1)
    add_para(
        doc,
        "附录主要说明可微光栅化器如何计算梯度。核心思想是：前向渲染中，像素颜色由按深度排序的高斯通过 alpha 合成得到；"
        "反向传播时，需要把最终图像损失对颜色的梯度，逐层传回每个高斯的颜色、透明度、二维均值和二维协方差，再进一步传到三维位置、尺度和旋转。"
    )
    add_para(
        doc,
        "因为透射率 T_i = Π_{j<i}(1-α_j) 会把前方高斯对后方高斯的遮挡关系编码进去，"
        "反向传播也必须考虑这种前后依赖。官方实现通过 CUDA 扩展完成这些计算，并把结果暴露给 PyTorch 的 autograd 系统。"
    )

    doc.add_heading("图表说明译文", level=1)
    add_para(
        doc,
        "图 1：本文方法以与最快先前方法相当的优化时间，实现与此前最高质量方法相当的辐射场实时渲染质量。"
        "关键在于新的三维高斯场景表示和实时可微渲染器，它们同时加速场景优化和新视角合成。"
    )
    add_para(
        doc,
        "其他图表：原文中的可视化结果主要展示不同方法在真实场景上的质量、速度和训练时间对比；"
        "表格则汇总 PSNR、SSIM、LPIPS、FPS 等指标。阅读时应同时关注质量指标和实时性指标，不能只看单一数值。"
    )

    doc.add_heading("与源码的对应关系", level=1)
    add_table(
        doc,
        ["论文概念", "源码位置/对象", "说明"],
        [
            ("三维高斯参数", "scene/gaussian_model.py: GaussianModel", "_xyz、_scaling、_rotation、_opacity、SH features。"),
            ("COLMAP 初始化", "create_from_pcd", "稀疏点云位置初始化为 Gaussian 中心，颜色转为 SH 初值。"),
            ("前向渲染", "gaussian_renderer/render", "打包相机和 Gaussian 参数并调用 CUDA rasterizer。"),
            ("图像损失", "train.py", "L1 + SSIM，必要时加入深度正则。"),
            ("反向传播", "loss.backward()", "调用自定义 rasterizer backward，把图像误差传回高斯参数。"),
            ("增密/剪枝", "densify_and_clone / densify_and_split / prune_points", "根据屏幕空间梯度、尺度和 opacity 改变 Gaussian 数量。"),
        ],
        [4.0, 5.0, 6.0],
    )

    doc.add_heading("参考文献说明", level=1)
    add_para(
        doc,
        "原论文参考文献较长，本文档不逐条翻译参考文献条目。阅读和引用时请以原 PDF 的 References 部分为准。"
        "如果需要，我可以后续单独生成一份“参考文献中文导读表”，按 NeRF、Mip-NeRF360、Instant-NGP、Plenoxels、点渲染、可微渲染等主题分类。"
    )

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
