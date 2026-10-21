# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\Users\Administrator\Documents\my 3dgs\3DGS_source_study_plan_compact_CN.docx")


def set_font(run, size=9.5, bold=False, color=None):
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    set_font(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_text(cell, header, bold=True)
        shade_cell(cell, "D9EAF7")
        if widths:
            cell.width = Cm(widths[i])

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
            if widths:
                cells[i].width = Cm(widths[i])

    doc.add_paragraph("")
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)

    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    doc.styles["Normal"].font.size = Pt(10.5)
    return doc


def main():
    doc = setup_document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("3DGS 源码入门学习计划（精简执行版）")
    set_font(run, size=18, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(
        "目标：不刷题、不泛学 Python，专门补足读懂 3DGS 官方源码所需的 Python / NumPy / PyTorch / 源码阅读能力。"
    )
    set_font(run, size=10, color=(80, 80, 80))

    doc.add_heading("一、学习原则", level=1)
    add_table(
        doc,
        ["项目", "要求"],
        [
            ("学习周期", "4 周；每天 1.5-2.5 小时；每周至少 1 次复盘。"),
            ("学习顺序", "Python 基础 -> NumPy/图像/矩阵 -> PyTorch 训练 -> 3DGS 源码主流程。"),
            ("当前不做", "不刷 LeetCode；不深入 CUDA；不逐行读完整仓库；不追复杂 Python 语法。"),
            ("最终能力", "能解释 Scene、Camera、GaussianModel、render、loss.backward、optimizer.step、densify/prune 的关系。"),
        ],
        [3.0, 13.5],
    )

    doc.add_heading("二、四周总路线", level=1)
    add_table(
        doc,
        ["周期", "主题", "重点", "验收标准"],
        [
            ("第 1 周", "Python 够用基础", "变量、list/dict、if/for、函数、class、import", "能看懂 train.py 中普通 Python 控制流。"),
            ("第 2 周", "NumPy、图像、矩阵", "array shape、mask、矩阵乘法、读图显示、画曲线", "能理解点云、图片、mask、协方差的数组形式。"),
            ("第 3 周", "PyTorch 最小训练", "tensor、requires_grad、backward、optimizer、参数组", "能理解 3DGS 的前向、loss、反向、更新。"),
            ("第 4 周", "回到 3DGS 源码", "train.py、GaussianModel、初始化、clone/split/prune", "能用自己的话讲清源码主流程。"),
        ],
        [2.0, 3.0, 6.0, 5.5],
    )

    weeks = [
        (
            "第 1 周：Python 够用基础",
            [
                ("Day 1", "变量与基本类型", "1.5h", "练 list/dict/print/type/len；用 dict 表示一个 Gaussian。", "能打印 xyz、opacity、scale。"),
                ("Day 2", "if / for / 函数", "2h", "写 should_prune(opacity)；遍历 opacity 列表。", "能判断 opacity 低的点该删除。"),
                ("Day 3", "list / dict 操作", "2h", "练 append/pop/items；模拟 viewpoint_stack 和 render_pkg。", "理解随机取 Camera、render_pkg['render']。"),
                ("Day 4", "class 入门", "2h", "写简化 Gaussian 类，包含 xyz 和 opacity。", "理解 self、属性、方法。"),
                ("Day 5", "import 与文件结构", "1.5h", "两个 .py 文件互相导入；对应 train.py 的 import。", "知道 train.py 调用其他模块。"),
                ("Day 6-7", "小项目：简化 prune", "3-4h", "创建 10 个 Gaussian，删除 opacity < 0.005 的项。", "能解释 prune 的基本思想。"),
            ],
        ),
        (
            "第 2 周：NumPy、图像和矩阵",
            [
                ("Day 8", "NumPy array 与 shape", "2h", "创建 points=[N,3]、colors=[N,3] 并打印 shape。", "理解点云数组形状。"),
                ("Day 9", "索引、切片、mask", "2h", "用 opacity mask 筛选 points。", "理解 prune_points 里的 mask。"),
                ("Day 10", "矩阵乘法", "2h", "用 S、R 计算 Sigma = R @ S @ S.T @ R.T。", "理解协方差构造公式。"),
                ("Day 11", "图片读写显示", "2h", "PIL 读图，np.array，plt.imshow。", "理解 [H,W,3] 和 [3,H,W]。"),
                ("Day 12", "画训练曲线", "1.5h", "用 matplotlib 画 loss 曲线并保存。", "能可视化训练变化。"),
                ("Day 13-14", "小项目：2D Gaussian 图像", "4-5h", "生成几个 2D Gaussian 软斑点并保存图片。", "理解 splatting 的二维版本。"),
            ],
        ),
        (
            "第 3 周：PyTorch 最小够用",
            [
                ("Day 15", "Tensor 基础", "2h", "练 torch.tensor、shape、requires_grad、cpu/cuda。", "知道 Tensor 支持自动求导。"),
                ("Day 16", "自动求导", "2h", "loss=(x-5)^2，backward，查看 x.grad。", "理解 backward 只算梯度。"),
                ("Day 17", "optimizer", "2h", "用 Adam 优化 x 接近 5。", "理解 step 更新参数，zero_grad 清梯度。"),
                ("Day 18", "多个参数组", "2h", "x/y 用不同学习率优化。", "理解 3DGS 不同参数不同 lr。"),
                ("Day 19", "拟合颜色", "2h", "随机 RGB 参数优化到目标颜色。", "类比 SH color 被 loss 优化。"),
                ("Day 20-21", "小项目：极简 2D splatting 训练", "5-6h", "用几个 2D Gaussian 拟合目标图。", "理解 render -> loss -> backward -> step。"),
            ],
        ),
        (
            "第 4 周：回到 3DGS 源码",
            [
                ("Day 22", "只看 train.py 主循环", "2h", "只找 Scene/GaussianModel/render/loss/backward/step。", "画出训练流程图。"),
                ("Day 23", "看 GaussianModel 参数", "2h", "对应 _xyz/_scaling/_rotation/_opacity/features。", "能把代码变量对应论文变量。"),
                ("Day 24", "看 create_from_pcd", "2h", "理解 COLMAP 点云如何初始化 Gaussian。", "知道 xyz/color/scale/rotation/opacity 来源。"),
                ("Day 25", "看 training_setup", "2h", "看 optimizer 参数组和学习率。", "知道哪些 Gaussian 参数被优化。"),
                ("Day 26", "看 clone 和 split", "2.5h", "读 densify_and_clone / densify_and_split。", "能解释“小的复制，大的拆分”。"),
                ("Day 27", "看 prune 和 opacity reset", "2h", "读 prune_points / densify_and_prune / reset_opacity。", "能解释删除条件和 opacity reset。"),
                ("Day 28", "总结源码一页纸", "2-3h", "回答数据、前向、loss、反向、增删、保存。", "能独立讲清 3DGS 训练主线。"),
            ],
        ),
    ]

    for title_text, rows in weeks:
        doc.add_heading(title_text, level=1)
        add_table(doc, ["时间", "小节", "建议时长", "学习任务", "完成标准"], rows, [1.8, 3.0, 2.0, 6.2, 4.5])

    doc.add_heading("三、每天固定执行模板", level=1)
    add_table(
        doc,
        ["时间", "动作"],
        [
            ("20 分钟", "看教程或源码讲解，只看当天主题。"),
            ("40 分钟", "自己敲一个最小例子，不复制长代码。"),
            ("20 分钟", "打印变量、shape、中间结果。"),
            ("20 分钟", "写三句话总结：今天学了什么；对应 3DGS 哪里；还没懂什么。"),
        ],
        [3.0, 13.5],
    )

    doc.add_heading("四、3DGS 源码对应速查", level=1)
    add_table(
        doc,
        ["源码对象/函数", "你需要记住的作用"],
        [
            ("Scene", "读取数据，组织 train/test cameras，初始化或加载 Gaussian。"),
            ("Camera", "保存真实图片 original_image 和相机参数 R/T/FoV/projection。"),
            ("GaussianModel", "保存所有可训练 Gaussian 参数：xyz、scale、rotation、opacity、SH features。"),
            ("render(...)", "把 Camera + Gaussian 参数交给 rasterizer，输出预测图。"),
            ("loss.backward()", "从图像误差反传梯度到 Gaussian 参数。"),
            ("optimizer.step()", "真正更新 Gaussian 参数。"),
            ("densify_and_clone", "梯度大且 Gaussian 较小：复制一个。"),
            ("densify_and_split", "梯度大且 Gaussian 较大：拆成多个小 Gaussian，并删掉原来的大 Gaussian。"),
            ("prune_points", "根据 mask 删除无用 Gaussian，同时同步 optimizer 状态。"),
            ("reset_opacity", "把 opacity 压低，避免错误遮挡过早固化。"),
        ],
        [4.2, 12.3],
    )

    doc.add_heading("五、最终自测清单", level=1)
    checks = [
        "我能解释图片如何变成 Camera.original_image。",
        "我能解释 COLMAP 点云如何初始化 _xyz。",
        "我能解释 _scaling、_opacity 为什么要经过 exp / sigmoid。",
        "我能说清一次训练迭代：选相机 -> render -> loss -> backward -> step。",
        "我能解释 viewspace gradient 为什么用于 densification。",
        "我能区分 clone、split、prune。",
        "我能写一个 50 行以内的小脚本分析数组 shape 或画 loss 曲线。",
    ]
    for check in checks:
        p = doc.add_paragraph()
        run = p.add_run("[ ] ")
        set_font(run, bold=True)
        run = p.add_run(check)
        set_font(run, size=10.5)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
